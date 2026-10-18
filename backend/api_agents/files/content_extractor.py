"""
File Content Extractor
----------------------
Extracts text content from various file types for agent context.
Supports: PDF, DOCX, TXT, code files, CSV, JSON, XML, HTML, Markdown
"""

import io
import traceback
from typing import Tuple, Optional


class FileContentExtractor:
    """Extract readable text from various file formats"""

    # Supported text-based MIME types
    TEXT_MIME_TYPES = {
        'text/plain',
        'text/html',
        'text/css',
        'text/javascript',
        'text/csv',
        'text/markdown',
        'text/xml',
        'application/json',
        'application/xml',
    }

    # Code file extensions
    CODE_EXTENSIONS = {
        '.py', '.js', '.ts', '.tsx', '.jsx', '.java', '.cpp', '.c', '.h',
        '.cs', '.php', '.rb', '.go', '.rs', '.swift', '.kt', '.scala',
        '.sh', '.bash', '.yml', '.yaml', '.toml', '.sql', '.r', '.m',
    }

    def extract(self, file_content: bytes, filename: str, mime_type: str) -> Tuple[Optional[str], str]:
        """
        Extract text content from file.

        Args:
            file_content: Binary file content
            filename: Original filename
            mime_type: MIME type of file

        Returns:
            Tuple of (extracted_text, extraction_status)
            status: 'success', 'failed', 'not_applicable'
        """
        try:
            # 1. PDF files
            if mime_type == 'application/pdf' or filename.lower().endswith('.pdf'):
                return self._extract_pdf(file_content)

            # 2. Microsoft Word (.docx)
            elif mime_type in ['application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                               'application/msword'] or filename.lower().endswith(('.docx', '.doc')):
                return self._extract_docx(file_content)

            # 3. Plain text and code files
            elif mime_type in self.TEXT_MIME_TYPES or any(filename.lower().endswith(ext) for ext in self.CODE_EXTENSIONS):
                return self._extract_text(file_content)

            # 4. Images, videos, audio - not applicable
            elif mime_type.startswith(('image/', 'video/', 'audio/')):
                return None, 'not_applicable'

            # 5. Binary files - not applicable
            else:
                return None, 'not_applicable'

        except Exception as e:
            print(f"❌ Content extraction error: {e}")
            traceback.print_exc()
            return None, 'failed'

    def _extract_pdf(self, file_content: bytes) -> Tuple[Optional[str], str]:
        """Extract text from PDF using PyPDF2"""
        try:
            import PyPDF2

            pdf_file = io.BytesIO(file_content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)

            # Extract text from all pages (max 50 pages to avoid huge files)
            extracted_text = []
            max_pages = min(len(pdf_reader.pages), 50)

            for page_num in range(max_pages):
                page = pdf_reader.pages[page_num]
                text = page.extract_text()
                if text:
                    extracted_text.append(text)

            if not extracted_text:
                return None, 'failed'

            full_text = '\n\n'.join(extracted_text)

            # Limit to 50,000 characters
            if len(full_text) > 50000:
                full_text = full_text[:50000] + "\n\n[Content truncated - file too large]"

            print(f"✅ Extracted {len(full_text)} characters from PDF ({max_pages} pages)")
            return full_text, 'success'

        except ImportError:
            print("⚠️ PyPDF2 not installed - skipping PDF extraction")
            return None, 'not_applicable'
        except Exception as e:
            print(f"❌ PDF extraction failed: {e}")
            return None, 'failed'

    def _extract_docx(self, file_content: bytes) -> Tuple[Optional[str], str]:
        """Extract text from DOCX using python-docx"""
        try:
            import docx

            docx_file = io.BytesIO(file_content)
            doc = docx.Document(docx_file)

            # Extract text from all paragraphs
            extracted_text = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    extracted_text.append(paragraph.text)

            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' | '.join(cell.text.strip() for cell in row.cells)
                    if row_text:
                        extracted_text.append(row_text)

            if not extracted_text:
                return None, 'failed'

            full_text = '\n\n'.join(extracted_text)

            # Limit to 50,000 characters
            if len(full_text) > 50000:
                full_text = full_text[:50000] + "\n\n[Content truncated - file too large]"

            print(f"✅ Extracted {len(full_text)} characters from DOCX")
            return full_text, 'success'

        except ImportError:
            print("⚠️ python-docx not installed - skipping DOCX extraction")
            return None, 'not_applicable'
        except Exception as e:
            print(f"❌ DOCX extraction failed: {e}")
            return None, 'failed'

    def _extract_text(self, file_content: bytes) -> Tuple[Optional[str], str]:
        """Extract text from plain text files"""
        try:
            # Try UTF-8 first
            try:
                text = file_content.decode('utf-8')
            except UnicodeDecodeError:
                # Fallback to latin-1 (always succeeds)
                text = file_content.decode('latin-1')

            # Limit to 50,000 characters
            if len(text) > 50000:
                text = text[:50000] + "\n\n[Content truncated - file too large]"

            print(f"✅ Extracted {len(text)} characters from text file")
            return text, 'success'

        except Exception as e:
            print(f"❌ Text extraction failed: {e}")
            return None, 'failed'


# Global instance
_content_extractor_instance = None


def get_content_extractor() -> FileContentExtractor:
    """Get or create the global content extractor instance"""
    global _content_extractor_instance
    if _content_extractor_instance is None:
        _content_extractor_instance = FileContentExtractor()
    return _content_extractor_instance
